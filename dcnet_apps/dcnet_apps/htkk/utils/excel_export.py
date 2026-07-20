"""
Excel Export utility for HTKK declarations.

Exports declaration data to Excel format with:
- Main sheet: Tờ khai chính với các chỉ tiêu
- Appendix sheets: Các phụ lục (nếu có)
"""

import frappe
from frappe.utils import flt, formatdate
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from dcnet_apps.htkk.declarations.base import get_declaration_config


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, size=14)
SUBHEADER_FONT = Font(bold=True, size=11)
NORMAL_FONT = Font(size=10)
BOLD_FONT = Font(bold=True, size=10)

CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SECTION_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")


# ---------------------------------------------------------------------------
# Main Export Function
# ---------------------------------------------------------------------------


def export_declaration_to_excel(declaration_id: str) -> bytes:
    """
    Xuất tờ khai ra file Excel.

    Args:
        declaration_id: ID của HTKK Declaration

    Returns:
        bytes: Nội dung file Excel

    Raises:
        frappe.DoesNotExistError: Nếu không tìm thấy tờ khai
    """
    doc = frappe.get_doc("HTKK Declaration", declaration_id)
    config = get_declaration_config(doc.declaration_type)

    if not config:
        frappe.throw(f"Không hỗ trợ loại tờ khai: {doc.declaration_type}")

    data = config.generate(doc)

    wb = openpyxl.Workbook()

    # Sheet 1: Tờ khai chính
    ws_main = wb.active
    ws_main.title = "Tờ khai"
    _write_main_sheet(ws_main, doc, data["chi_tieu"], data.get("sources", {}), config)

    # Sheet 2+: Phụ lục (nếu có)
    phu_luc = data.get("phu_luc", {})
    for appendix_key, appendix_data in phu_luc.items():
        if appendix_data:
            # Excel sheet name max 31 chars
            sheet_name = appendix_key[:31]
            ws_app = wb.create_sheet(title=sheet_name)
            _write_appendix_sheet(ws_app, appendix_key, appendix_data)

    # Save to bytes
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# ---------------------------------------------------------------------------
# Main Sheet
# ---------------------------------------------------------------------------


def _write_main_sheet(ws, doc, chi_tieu, sources, config):
    """Write main declaration sheet."""
    row = 1

    # --- Header ---
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value=config.title)
    cell.font = HEADER_FONT
    cell.alignment = CENTER_ALIGN
    row += 1

    if config.subtitle:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=config.subtitle)
        cell.font = NORMAL_FONT
        cell.alignment = CENTER_ALIGN
        row += 1

    row += 1  # Empty row

    # --- Company Info ---
    info_items = [
        ("Công ty:", doc.company),
        ("Mã số thuế:", frappe.db.get_value("Company", doc.company, "tax_id") or ""),
        ("Kỳ kê khai:", _format_period(doc)),
        ("Từ ngày:", formatdate(doc.from_date, "dd/MM/yyyy")),
        ("Đến ngày:", formatdate(doc.to_date, "dd/MM/yyyy")),
    ]

    for label, value in info_items:
        ws.cell(row=row, column=1, value=label).font = BOLD_FONT
        ws.cell(row=row, column=2, value=value).font = NORMAL_FONT
        row += 1

    row += 1  # Empty row

    # --- Table Header ---
    headers = ["Mã chỉ tiêu", "Nội dung", "Giá trị", "Nguồn dữ liệu"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = CENTER_ALIGN
    row += 1

    # --- Indicators by Section ---
    for section in config.sections:
        # Section header
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=section.get("title", ""))
        cell.font = BOLD_FONT
        cell.fill = SECTION_FILL
        cell.border = THIN_BORDER
        row += 1

        # Indicators in section
        for ct_name in section.get("indicators", []):
            value = chi_tieu.get(ct_name, 0)
            source = sources.get(ct_name, "")
            label = _get_indicator_label(ct_name, config)

            ws.cell(row=row, column=1, value=ct_name).border = THIN_BORDER
            ws.cell(row=row, column=2, value=label).border = THIN_BORDER
            cell_value = ws.cell(row=row, column=3, value=_format_value(value))
            cell_value.border = THIN_BORDER
            cell_value.alignment = RIGHT_ALIGN
            cell_value.number_format = "#,##0"
            ws.cell(row=row, column=4, value=source).border = THIN_BORDER
            row += 1

        # Subsections
        for subsection in section.get("subsections", []):
            # Subsection header
            indent = "  " * subsection.get("indent", 1)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            cell = ws.cell(row=row, column=1, value=f"{indent}{subsection.get('title', '')}")
            cell.font = SUBHEADER_FONT
            cell.border = THIN_BORDER
            row += 1

            # Indicators in subsection
            for ct_name in subsection.get("indicators", []):
                value = chi_tieu.get(ct_name, 0)
                source = sources.get(ct_name, "")
                label = _get_indicator_label(ct_name, config)

                ws.cell(row=row, column=1, value=ct_name).border = THIN_BORDER
                ws.cell(row=row, column=2, value=f"{indent}{label}").border = THIN_BORDER
                cell_value = ws.cell(row=row, column=3, value=_format_value(value))
                cell_value.border = THIN_BORDER
                cell_value.alignment = RIGHT_ALIGN
                cell_value.number_format = "#,##0"
                ws.cell(row=row, column=4, value=source).border = THIN_BORDER
                row += 1

    # --- Set column widths ---
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 40


# ---------------------------------------------------------------------------
# Appendix Sheet
# ---------------------------------------------------------------------------


def _write_appendix_sheet(ws, appendix_key, data):
    """Write appendix sheet with tabular data."""
    if not data:
        ws.cell(row=1, column=1, value="Không có dữ liệu")
        return

    row = 1

    # Sheet title
    ws.cell(row=row, column=1, value=f"Phụ lục: {appendix_key}")
    ws.cell(row=row, column=1).font = HEADER_FONT
    row += 2

    # Get headers from first item
    if isinstance(data, list) and len(data) > 0:
        first_item = data[0]
        if isinstance(first_item, dict):
            headers = list(first_item.keys())

            # Write headers
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = HEADER_FILL
                cell.border = THIN_BORDER
                cell.alignment = CENTER_ALIGN
            row += 1

            # Write data rows
            for item in data:
                for col, key in enumerate(headers, 1):
                    value = item.get(key, "")
                    cell = ws.cell(row=row, column=col, value=_format_value(value))
                    cell.border = THIN_BORDER
                    if isinstance(value, (int, float)):
                        cell.alignment = RIGHT_ALIGN
                        cell.number_format = "#,##0"
                row += 1

            # Auto-adjust column widths
            for col_idx, header in enumerate(headers, 1):
                column_letter = get_column_letter(col_idx)
                max_length = len(str(header))
                for item in data:
                    cell_value = str(item.get(header, ""))
                    if len(cell_value) > max_length:
                        max_length = len(cell_value)
                ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _format_period(doc):
    """Format period for display."""
    period_type = doc.period_type
    if period_type == "Tháng":
        return f"Tháng {doc.period}/{doc.year}"
    elif period_type == "Quý":
        return f"Quý {doc.period}/{doc.year}"
    elif period_type == "Năm":
        return f"Năm {doc.year}"
    return f"{doc.period_type} {doc.period}/{doc.year}"


def _format_value(value):
    """Format value for Excel cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, (int, float)):
        return flt(value)
    return str(value)


def _get_indicator_label(ct_name, config):
    """Get label for indicator from config or module."""
    # Try to get from config's module CT_LABELS
    try:
        import importlib
        module_name = config.__class__.__module__
        module = importlib.import_module(module_name)
        if hasattr(module, "CT_LABELS"):
            labels = getattr(module, "CT_LABELS")
            if ct_name in labels:
                return labels[ct_name]
    except Exception:
        pass

    return ct_name
