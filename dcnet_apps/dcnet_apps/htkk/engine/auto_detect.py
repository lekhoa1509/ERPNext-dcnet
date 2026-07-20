
import re
import openpyxl
import xlrd
import os
from typing import List, Dict, Any

class IndicatorMatch:
    def __init__(self, code: str, cell: str, confidence: float, anchor_text: str, anchor_cell: str):
        self.code = code
        self.cell = cell
        self.confidence = confidence
        self.anchor_text = anchor_text
        self.anchor_cell = anchor_cell

    def to_dict(self):
        return {
            "code": self.code,
            "cell": self.cell,
            "confidence": self.confidence,
            "anchor_text": self.anchor_text,
            "anchor_cell": self.anchor_cell
        }

def get_column_letter(n):
    """Chuyển số cột sang chữ cái (0 -> A, 1 -> B, ...)"""
    string = ""
    n += 1
    while n > 0:
        n, remainder = divmod(n - 1, 26)
        string = chr(65 + remainder) + string
    return string

def address_from_row_col(row, col):
    """Chuyển row, col sang format A1"""
    return f"{get_column_letter(col)}{row + 1}"

def auto_detect_indicators(file_path: str) -> List[Dict[str, Any]]:
    ext = os.path.splitext(file_path)[1].lower()
    results = []
    if ext == '.xlsx':
        results = scan_xlsx(file_path)
    elif ext == '.xls':
        results = scan_xls(file_path)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
    return results

def scan_xlsx(file_path: str):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    results = []
    indicator_pattern = re.compile(r'\[([0-9a-zA-Z\.]+)\]')
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not cell.value or not isinstance(cell.value, str):
                    continue
                match = indicator_pattern.search(cell.value)
                if match:
                    code = match.group(0)
                    target_cell, confidence = find_target_cell_xlsx(sheet, cell)
                    results.append(IndicatorMatch(
                        code=code,
                        cell=target_cell,
                        confidence=confidence,
                        anchor_text=cell.value,
                        anchor_cell=cell.coordinate
                    ).to_dict())
    return results

def scan_xls(file_path: str):
    wb = xlrd.open_workbook(file_path)
    results = []
    indicator_pattern = re.compile(r'\[([0-9a-zA-Z\.]+)\]')
    for sheet in wb.sheets():
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                val = sheet.cell_value(r, c)
                if not val or not isinstance(val, str):
                    continue
                match = indicator_pattern.search(val)
                if match:
                    code = match.group(0)
                    target_cell, confidence = find_target_cell_xls(sheet, r, c)
                    results.append(IndicatorMatch(
                        code=code,
                        cell=target_cell,
                        confidence=confidence,
                        anchor_text=val,
                        anchor_cell=address_from_row_col(r, c)
                    ).to_dict())
    return results

def find_target_cell_xlsx(sheet, anchor_cell):
    row = anchor_cell.row
    col = anchor_cell.column
    try:
        right = sheet.cell(row=row, column=col+1)
        if right.value is None or right.value == 0:
            return right.coordinate, 0.95
    except: pass
    try:
        down = sheet.cell(row=row+1, column=col)
        if down.value is None or down.value == 0:
            return down.coordinate, 0.75
    except: pass
    return anchor_cell.coordinate, 0.4

def find_target_cell_xls(sheet, r, c):
    if c + 1 < sheet.ncols:
        val = sheet.cell_value(r, c+1)
        if val == "" or val == 0:
            return address_from_row_col(r, c+1), 0.95
    if r + 1 < sheet.nrows:
        val = sheet.cell_value(r+1, c)
        if val == "" or val == 0:
            return address_from_row_col(r+1, c), 0.75
    return address_from_row_col(r, c), 0.4
