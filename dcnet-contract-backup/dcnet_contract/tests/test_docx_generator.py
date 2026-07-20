"""Tests for docx generator engine."""
import os
import sys
import tempfile
import unittest

# Add app to path
sys.path.insert(0, "/home/long/long/frappe-bench-dcnet/apps/dcnet_contract")

from dcnet_contract.dcnet_contract.utils.docx_generator import (
    _number_to_words_vi,
    _format_date_vi,
    _format_currency,
    _merge_runs_for_placeholders,
    _replace_in_paragraph,
    discover_placeholders,
    _fill_docx,
    _parse_currency,
)


class TestNumberToWordsVi(unittest.TestCase):
    def test_zero(self):
        self.assertEqual(_number_to_words_vi(0), "Khong dong")

    def test_one_thousand(self):
        result = _number_to_words_vi(1000)
        self.assertIn("nghin", result.lower())
        self.assertTrue(result.endswith("dong"))

    def test_one_million(self):
        result = _number_to_words_vi(1000000)
        self.assertIn("trieu", result.lower())
        self.assertTrue(result.endswith("dong"))

    def test_3_300_000(self):
        result = _number_to_words_vi(3300000)
        self.assertIn("trieu", result.lower())
        self.assertIn("nghin", result.lower())
        self.assertTrue(result.endswith("dong"))

    def test_15_450_000(self):
        result = _number_to_words_vi(15450000)
        self.assertIn("trieu", result.lower())
        self.assertTrue(result.endswith("dong"))

    def test_100(self):
        result = _number_to_words_vi(100)
        self.assertIn("tram", result.lower())

    def test_large_number(self):
        result = _number_to_words_vi(1500000000)  # 1.5 ty
        self.assertIn("ty", result.lower())
        self.assertTrue(result.endswith("dong"))


class TestFormatDateVi(unittest.TestCase):
    def test_none(self):
        self.assertEqual(_format_date_vi(None), "")

    def test_valid_date(self):
        from datetime import date
        result = _format_date_vi(date(2026, 4, 15))
        self.assertEqual(result, "ngay 15 thang 04 nam 2026")


class TestFormatCurrency(unittest.TestCase):
    def test_zero(self):
        self.assertEqual(_format_currency(0), "0")

    def test_millions(self):
        self.assertEqual(_format_currency(3000000), "3,000,000")

    def test_with_hundreds(self):
        self.assertEqual(_format_currency(15450000), "15,450,000")


class TestParseCurrency(unittest.TestCase):
    def test_with_commas(self):
        self.assertEqual(_parse_currency("3,000,000"), 3000000)

    def test_empty(self):
        self.assertEqual(_parse_currency(""), 0)


class TestFillDocx(unittest.TestCase):
    def test_simple_replacement(self):
        """Create a .docx with 2 placeholders, verify they're replaced."""
        from docx import Document

        # Create temp template
        doc = Document()
        doc.add_paragraph("Contract: {{contract_number}}")
        doc.add_paragraph("Customer: {{customer_name}}")

        tmp_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp_in.name)
        tmp_in.close()

        values = {"contract_number": "HĐ-001", "customer_name": "Cong ty ABC"}
        result_path = _fill_docx(tmp_in.name, values, [])

        # Verify
        result_doc = Document(result_path)
        texts = [p.text for p in result_doc.paragraphs]
        full_text = " ".join(texts)
        self.assertIn("HĐ-001", full_text)
        self.assertIn("Cong ty ABC", full_text)
        self.assertNotIn("{{", full_text)

        os.unlink(tmp_in.name)
        os.unlink(result_path)

    def test_items_table_clone(self):
        """Create .docx with item table template row, verify N rows created."""
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = "STT"
        table.rows[0].cells[1].text = "Hang muc"
        table.rows[0].cells[2].text = "Thanh tien"
        table.rows[1].cells[0].text = "{{item_stt}}"
        table.rows[1].cells[1].text = "{{item_label}}"
        table.rows[1].cells[2].text = "{{item_amount}}"

        tmp_in = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp_in.name)
        tmp_in.close()

        items = [
            {"item_stt": "1", "item_label": "Kenh P2P", "item_qty": "1", "item_uom": "kenh", "item_price": "5,000,000", "item_amount": "5,000,000"},
            {"item_stt": "2", "item_label": "Kenh ILL", "item_qty": "2", "item_uom": "kenh", "item_price": "3,000,000", "item_amount": "6,000,000"},
        ]
        result_path = _fill_docx(tmp_in.name, {}, items)

        result_doc = Document(result_path)
        # Should have: header row + 2 item rows + 3 summary rows = 6 rows
        result_table = result_doc.tables[0]
        # Template row removed, so: 1 header + 2 items + 3 summary = 6
        self.assertEqual(len(result_table.rows), 6)

        # Check item data is present
        row1_text = " ".join(cell.text for cell in result_table.rows[1].cells)
        self.assertIn("Kenh P2P", row1_text)
        row2_text = " ".join(cell.text for cell in result_table.rows[2].cells)
        self.assertIn("Kenh ILL", row2_text)

        os.unlink(tmp_in.name)
        os.unlink(result_path)


class TestRunMerging(unittest.TestCase):
    def test_split_placeholder_merges(self):
        """Verify {{place + holder}} across 2 runs merges correctly."""
        from docx import Document

        doc = Document()
        p = doc.add_paragraph()
        run1 = p.add_run("Hello {{contract")
        run2 = p.add_run("_number}} world")

        _merge_runs_for_placeholders(p)

        # After merge, first run should have full text
        full = p.runs[0].text
        self.assertIn("{{contract_number}}", full)


class TestDiscoverPlaceholders(unittest.TestCase):
    def test_discover(self):
        """Verify placeholder discovery from a .docx file."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("Name: {{customer_name}}")
        doc.add_paragraph("Date: {{contract_date}}")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "{{item_stt}}"
        table.rows[0].cells[1].text = "{{item_label}}"

        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp.name)
        tmp.close()

        results = discover_placeholders(tmp.name)
        keys = [r["placeholder_key"] for r in results]
        self.assertIn("customer_name", keys)
        self.assertIn("contract_date", keys)
        self.assertIn("item_stt", keys)
        self.assertIn("item_label", keys)

        os.unlink(tmp.name)


class TestPdfConversion(unittest.TestCase):
    def test_libreoffice_available(self):
        """Check if LibreOffice is available for PDF conversion."""
        import subprocess
        result = subprocess.run(["libreoffice", "--version"], capture_output=True)
        if result.returncode != 0:
            self.skipTest("LibreOffice not available")

        from docx import Document
        doc = Document()
        doc.add_paragraph("Test PDF")
        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        doc.save(tmp.name)
        tmp.close()

        from dcnet_contract.dcnet_contract.utils.docx_generator import _convert_to_pdf
        pdf_path = _convert_to_pdf(tmp.name)
        self.assertTrue(os.path.exists(pdf_path))
        self.assertTrue(pdf_path.endswith(".pdf"))

        os.unlink(tmp.name)
        os.unlink(pdf_path)


if __name__ == "__main__":
    # Run with verbose output
    loader = unittest.TestLoader()
    suite = loader.loadTestsFromModule(sys.modules[__name__])
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    sys.exit(0 if result.wasSuccessful() else 1)
