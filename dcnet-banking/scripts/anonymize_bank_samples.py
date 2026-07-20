"""Anonymize DCNet bank statement samples for committing to tests/fixtures/.

Reads source files from `docs/accounting-requirements/banking/` (bench-level),
replaces company names / account numbers / customer IDs with TEST values,
and writes anonymized copies into the app's test fixtures directory.

Usage from bench root:
    ./env/bin/python apps/vn_banking/scripts/anonymize_bank_samples.py
"""
from pathlib import Path

BENCH_ROOT = Path(__file__).resolve().parents[3]  # apps/vn_banking/scripts → bench root
SRC = BENCH_ROOT / "docs" / "accounting-requirements" / "banking"
DST = BENCH_ROOT / "apps" / "vn_banking" / "vn_banking" / "tests" / "fixtures"
DST.mkdir(parents=True, exist_ok=True)

REPLACEMENTS = {
    "CONG TY CO PHAN VIEN THONG DCNET": "TEST COMPANY",
    "CTY CP VIEN THONG DCNET": "TEST COMPANY",
    "CTCP VIEN THONG DCNET": "TEST COMPANY",
    "8414237": "TEST0001",
    "01436599": "TEST0002",
    "0312976157": "0000000000",
    "262086313": "1111111111",
    "6555555": "2222222222",
    "6565655555": "3333333333",
    "919299999": "4444444444",
    "137 LE QUANG DINH Phuong 14 Quan Binh Tha": "TEST ADDRESS",
    "DCNET TELECOM": "TEST TELECOM",
    "DCNET": "TESTCO",
}

# Source filename → target filename (all .xls binary stay .xls; html-as-xls stays .xls)
COPIES = {
    "AccountStmt_262086313.xls": ("bidv_sample.xls", "xls_binary"),
    "Bang ket qua giao dich - 6565655555.xlsx": ("mb_sample.xlsx", "xlsx"),
    "Sao_Ke_Giao_Dich_919299999.xls": ("sacombank_sample.xls", "xls_binary"),
    "TranDetail_pgb.xls": ("pgbank_sample.xls", "html"),
}


def _replace_all(text: str) -> str:
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    return text


def anonymize_xlsx(src: Path, dst: Path):
    from openpyxl import load_workbook
    wb = load_workbook(src, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _replace_all(cell.value)
    wb.save(dst)


def anonymize_xls_binary(src: Path, dst: Path):
    import xlrd
    import xlwt
    rb = xlrd.open_workbook(str(src))
    wb = xlwt.Workbook()
    for si in range(rb.nsheets):
        rs = rb.sheet_by_index(si)
        ws = wb.add_sheet(rs.name or f"Sheet{si}")
        for r in range(rs.nrows):
            for c in range(rs.ncols):
                val = rs.cell_value(r, c)
                if isinstance(val, str):
                    val = _replace_all(val)
                ws.write(r, c, val)
    wb.save(str(dst))


def anonymize_html(src: Path, dst: Path):
    raw = src.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    dst.write_bytes(_replace_all(text).encode("utf-8"))


def main():
    for src_name, (dst_name, kind) in COPIES.items():
        src = SRC / src_name
        dst = DST / dst_name
        if not src.exists():
            print(f"MISSING: {src} — skip")
            continue
        if kind == "xlsx":
            anonymize_xlsx(src, dst)
        elif kind == "xls_binary":
            anonymize_xls_binary(src, dst)
        elif kind == "html":
            anonymize_html(src, dst)
        print(f"OK: {src_name} → tests/fixtures/{dst_name}")


if __name__ == "__main__":
    main()
