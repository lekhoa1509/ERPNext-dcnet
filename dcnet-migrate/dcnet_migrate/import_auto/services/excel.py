import csv
import io
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import random

import frappe
import xlrd
from openpyxl import load_workbook

from dcnet_migrate.import_auto.services.utils import normalize_key


CSV_EXTENSIONS = {".csv"}
CSV_DELIMITERS = [",", ";", "\t", "|"]


MAX_COLUMNS = 80
MAX_SCAN_ROWS = 80
AI_SAMPLE_ROWS = 10


def summarize_workbook(
    file_path: str,
    max_sample_rows: int = 20,
    sample_strategy: str = "first",
    sample_seed: str | None = None,
) -> dict:
    suffix = Path(file_path).suffix.lower()
    if suffix in CSV_EXTENSIONS:
        return _summarize_csv_workbook(file_path, max_sample_rows, sample_strategy, sample_seed)
    if suffix == ".xls":
        return _summarize_xls_workbook(file_path, max_sample_rows, sample_strategy, sample_seed)
    return _summarize_xlsx_workbook(file_path, max_sample_rows, sample_strategy, sample_seed)


def extract_records(file_path: str, sheet_name: str | None, header_row_number: int | None) -> list[dict]:
    suffix = Path(file_path).suffix.lower()
    if suffix in CSV_EXTENSIONS:
        return _extract_csv_records(file_path, header_row_number)
    if suffix == ".xls":
        return _extract_xls_records(file_path, sheet_name, header_row_number)
    return _extract_xlsx_records(file_path, sheet_name, header_row_number)


def sample_records(
    file_path: str,
    sheet_name: str | None,
    header_row_number: int | None,
    sample_size: int = AI_SAMPLE_ROWS,
    seed: str | None = None,
) -> list[dict]:
    suffix = Path(file_path).suffix.lower()
    if suffix in CSV_EXTENSIONS:
        return _sample_csv_records(file_path, header_row_number, sample_size, seed)
    if suffix == ".xls":
        return _sample_xls_records(file_path, sheet_name, header_row_number, sample_size, seed)
    return _sample_xlsx_records(file_path, sheet_name, header_row_number, sample_size, seed)


def detect_header_row(rows: list[dict]) -> dict | None:
    best_row = None
    best_score = -1

    for index, row in enumerate(rows[:MAX_SCAN_ROWS]):
        values = [frappe.as_unicode(value or "").strip() for value in row["values"]]
        non_empty = [value for value in values if value]
        if len(non_empty) < 2:
            continue

        normalized = [normalize_key(value) for value in non_empty]
        unique_ratio = len(set(normalized)) / len(non_empty)
        has_stt = any(value in {"stt", "số thứ tự"} for value in normalized)
        text_count = sum(1 for value in non_empty if not _looks_numeric(value))
        next_non_empty = 0
        if index + 1 < len(rows):
            next_non_empty = sum(1 for value in rows[index + 1]["values"] if value not in (None, ""))

        score = len(non_empty) + text_count + (unique_ratio * 5) + min(next_non_empty, len(non_empty))
        if has_stt:
            score += 10

        if score > best_score:
            best_row = {"row_number": row["row_number"], "values": values}
            best_score = score

    return best_row


def _summarize_xlsx_workbook(
    file_path: str,
    max_sample_rows: int,
    sample_strategy: str,
    sample_seed: str | None,
) -> dict:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            preview_rows = []
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [_to_jsonable(value) for value in row[:MAX_COLUMNS]]
                if any(value not in (None, "") for value in values):
                    preview_rows.append({"row_number": row_number, "values": values})
                if len(preview_rows) >= MAX_SCAN_ROWS:
                    break

            header = detect_header_row(preview_rows)
            sampled_rows = None
            if _use_random_sample(sample_strategy, max_sample_rows) and header:
                sampled_rows = _sample_xlsx_records_from_worksheet(
                    worksheet,
                    header["row_number"],
                    header["values"],
                    max_sample_rows,
                    _sheet_sample_seed(file_path, worksheet.title, sample_seed),
                )
            sheets.append(
                _sheet_summary(
                    worksheet.title,
                    worksheet.max_row,
                    worksheet.max_column,
                    preview_rows,
                    header,
                    max_sample_rows,
                    sample_rows=sampled_rows,
                    sample_strategy=sample_strategy,
                )
            )
        return {"file_name": Path(file_path).name, "file_path": file_path, "sheets": sheets}
    finally:
        workbook.close()


def _summarize_xls_workbook(
    file_path: str,
    max_sample_rows: int,
    sample_strategy: str,
    sample_seed: str | None,
) -> dict:
    workbook = xlrd.open_workbook(file_path)
    sheets = []
    for worksheet in workbook.sheets():
        preview_rows = []
        for row_index in range(worksheet.nrows):
            values = [_to_jsonable(value) for value in worksheet.row_values(row_index)[:MAX_COLUMNS]]
            if any(value not in (None, "") for value in values):
                preview_rows.append({"row_number": row_index + 1, "values": values})
            if len(preview_rows) >= MAX_SCAN_ROWS:
                break

        header = detect_header_row(preview_rows)
        sampled_rows = None
        if _use_random_sample(sample_strategy, max_sample_rows) and header:
            sampled_rows = _sample_xls_records_from_worksheet(
                worksheet,
                header["row_number"],
                header["values"],
                max_sample_rows,
                _sheet_sample_seed(file_path, worksheet.name, sample_seed),
            )
        sheets.append(
            _sheet_summary(
                worksheet.name,
                worksheet.nrows,
                worksheet.ncols,
                preview_rows,
                header,
                max_sample_rows,
                sample_rows=sampled_rows,
                sample_strategy=sample_strategy,
            )
        )
    return {"file_name": Path(file_path).name, "file_path": file_path, "sheets": sheets}


def _sheet_summary(
    sheet_name: str,
    max_row: int,
    max_column: int,
    preview_rows: list[dict],
    header: dict | None,
    max_sample_rows: int,
    sample_rows: list[dict] | None = None,
    sample_strategy: str = "first",
) -> dict:
    sample_rows = sample_rows or []
    if header and not sample_rows:
        for row in preview_rows:
            if row["row_number"] <= header["row_number"]:
                continue
            if len(sample_rows) >= max_sample_rows:
                break
            sample_rows.append(_row_to_dict(header["values"], row["values"]))

    return {
        "sheet_name": sheet_name,
        "max_row": max_row,
        "max_column": max_column,
        "detected_header_row": header["row_number"] if header else None,
        "detected_headers": header["values"] if header else [],
        "sample_rows": sample_rows,
        "sample_strategy": "random" if _use_random_sample(sample_strategy, max_sample_rows) else "first",
        "sample_row_count": len(sample_rows),
    }


def _extract_xlsx_records(file_path: str, sheet_name: str | None, header_row_number: int | None) -> list[dict]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        header_row_number = header_row_number or 1
        header_values = None
        records = []
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = [_to_jsonable(value) for value in row[:MAX_COLUMNS]]
            if row_number == header_row_number:
                header_values = [frappe.as_unicode(value or "").strip() for value in values]
                continue
            if not header_values or row_number <= header_row_number:
                continue
            if not any(value not in (None, "") for value in values):
                continue
            records.append(_row_to_dict(header_values, values))
        return records
    finally:
        workbook.close()


def _extract_xls_records(file_path: str, sheet_name: str | None, header_row_number: int | None) -> list[dict]:
    workbook = xlrd.open_workbook(file_path)
    worksheet = workbook.sheet_by_name(sheet_name) if sheet_name and sheet_name in workbook.sheet_names() else workbook.sheet_by_index(0)
    header_row_number = header_row_number or 1
    if worksheet.nrows < header_row_number:
        return []

    header_values = [frappe.as_unicode(value or "").strip() for value in worksheet.row_values(header_row_number - 1)[:MAX_COLUMNS]]
    records = []
    for row_index in range(header_row_number, worksheet.nrows):
        values = [_to_jsonable(value) for value in worksheet.row_values(row_index)[:MAX_COLUMNS]]
        if not any(value not in (None, "") for value in values):
            continue
        records.append(_row_to_dict(header_values, values))
    return records


def _sample_xlsx_records(
    file_path: str,
    sheet_name: str | None,
    header_row_number: int | None,
    sample_size: int,
    seed: str | None,
) -> list[dict]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        header_row_number = header_row_number or 1
        header_values = None
        for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number == header_row_number:
                values = [_to_jsonable(value) for value in row[:MAX_COLUMNS]]
                header_values = [frappe.as_unicode(value or "").strip() for value in values]
                break
        if not header_values:
            return []
        return _sample_xlsx_records_from_worksheet(
            worksheet,
            header_row_number,
            header_values,
            sample_size,
            _sheet_sample_seed(file_path, worksheet.title, seed),
        )
    finally:
        workbook.close()


def _sample_xlsx_records_from_worksheet(
    worksheet,
    header_row_number: int,
    header_values: list,
    sample_size: int,
    seed: str | None,
) -> list[dict]:
    rng = random.Random(seed or "import-auto")
    reservoir: list[tuple[int, dict]] = []
    seen = 0

    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number <= header_row_number:
            continue
        values = [_to_jsonable(value) for value in row[:MAX_COLUMNS]]
        if not _has_values(values):
            continue
        seen += 1
        record = _row_to_dict(header_values, values)
        _reservoir_add(reservoir, (row_number, record), seen, sample_size, rng)

    return [record for _row_number, record in sorted(reservoir, key=lambda item: item[0])]


def _sample_xls_records(
    file_path: str,
    sheet_name: str | None,
    header_row_number: int | None,
    sample_size: int,
    seed: str | None,
) -> list[dict]:
    workbook = xlrd.open_workbook(file_path)
    worksheet = workbook.sheet_by_name(sheet_name) if sheet_name and sheet_name in workbook.sheet_names() else workbook.sheet_by_index(0)
    header_row_number = header_row_number or 1
    if worksheet.nrows < header_row_number:
        return []
    header_values = [frappe.as_unicode(value or "").strip() for value in worksheet.row_values(header_row_number - 1)[:MAX_COLUMNS]]
    return _sample_xls_records_from_worksheet(
        worksheet,
        header_row_number,
        header_values,
        sample_size,
        _sheet_sample_seed(file_path, worksheet.name, seed),
    )


def _sample_xls_records_from_worksheet(
    worksheet,
    header_row_number: int,
    header_values: list,
    sample_size: int,
    seed: str | None,
) -> list[dict]:
    rng = random.Random(seed or "import-auto")
    reservoir: list[tuple[int, dict]] = []
    seen = 0

    for row_index in range(header_row_number, worksheet.nrows):
        values = [_to_jsonable(value) for value in worksheet.row_values(row_index)[:MAX_COLUMNS]]
        if not _has_values(values):
            continue
        seen += 1
        record = _row_to_dict(header_values, values)
        _reservoir_add(reservoir, (row_index + 1, record), seen, sample_size, rng)

    return [record for _row_number, record in sorted(reservoir, key=lambda item: item[0])]


def _detect_csv_dialect(raw_bytes: bytes) -> tuple[str, str]:
    """Detect encoding (handles BOM) and delimiter for a CSV file."""
    if raw_bytes.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    else:
        encoding = "utf-8"

    sample = raw_bytes[:4096].decode(encoding, errors="replace")
    counts = {delimiter: sample.count(delimiter) for delimiter in CSV_DELIMITERS}
    delimiter = max(counts, key=lambda d: counts[d]) if any(counts.values()) else ","
    return encoding, delimiter


def _read_csv_rows(file_path: str, max_rows: int | None = None) -> tuple[list[list], str]:
    """Read CSV rows as lists of strings. Returns (rows, sheet_name)."""
    raw = Path(file_path).read_bytes()
    encoding, delimiter = _detect_csv_dialect(raw)
    text = raw.decode(encoding, errors="replace")

    sheet_name = Path(file_path).stem
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for row in reader:
        rows.append([cell.strip() for cell in row])
        if max_rows and len(rows) >= max_rows:
            break
    return rows, sheet_name


def _summarize_csv_workbook(
    file_path: str,
    max_sample_rows: int,
    sample_strategy: str,
    sample_seed: str | None,
) -> dict:
    # Single read: parse preview rows AND count total in one pass.
    raw = Path(file_path).read_bytes()
    encoding, delimiter = _detect_csv_dialect(raw)
    text = raw.decode(encoding, errors="replace")
    sheet_name = Path(file_path).stem

    preview_rows = []
    total_rows = 0
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for row_index, row in enumerate(reader):
        values = [_to_jsonable(cell.strip()) for cell in row[:MAX_COLUMNS]]
        if not any(v not in (None, "") for v in values):
            continue
        total_rows += 1
        if len(preview_rows) < MAX_SCAN_ROWS:
            preview_rows.append({"row_number": row_index + 1, "values": values})

    header = detect_header_row(preview_rows)
    sampled_rows = None
    if _use_random_sample(sample_strategy, max_sample_rows) and header:
        # _sample_csv_records reads the file once more; acceptable for random sampling.
        sampled_rows = _sample_csv_records(
            file_path,
            header["row_number"],
            max_sample_rows,
            _sheet_sample_seed(file_path, sheet_name, sample_seed),
        )

    sheet = _sheet_summary(
        sheet_name,
        total_rows,
        len(header["values"]) if header else 0,
        preview_rows,
        header,
        max_sample_rows,
        sample_rows=sampled_rows,
        sample_strategy=sample_strategy,
    )
    return {"file_name": Path(file_path).name, "file_path": file_path, "sheets": [sheet]}


def _extract_csv_records(file_path: str, header_row_number: int | None) -> list[dict]:
    rows, _sheet_name = _read_csv_rows(file_path)
    header_row_number = header_row_number or 1
    if len(rows) < header_row_number:
        return []

    header_values = [frappe.as_unicode(cell or "").strip() for cell in rows[header_row_number - 1][:MAX_COLUMNS]]
    records = []
    for row in rows[header_row_number:]:
        values = [_to_jsonable(cell) for cell in row[:MAX_COLUMNS]]
        if not any(v not in (None, "") for v in values):
            continue
        records.append(_row_to_dict(header_values, values))
    return records


def _sample_csv_records(
    file_path: str,
    header_row_number: int | None,
    sample_size: int,
    seed: str | None,
) -> list[dict]:
    rows, sheet_name = _read_csv_rows(file_path)
    header_row_number = header_row_number or 1
    if len(rows) < header_row_number:
        return []

    header_values = [frappe.as_unicode(cell or "").strip() for cell in rows[header_row_number - 1][:MAX_COLUMNS]]
    rng = random.Random(_sheet_sample_seed(file_path, sheet_name, seed))
    reservoir: list[tuple[int, dict]] = []
    seen = 0

    for row_index, row in enumerate(rows[header_row_number:], start=header_row_number + 1):
        values = [_to_jsonable(cell) for cell in row[:MAX_COLUMNS]]
        if not _has_values(values):
            continue
        seen += 1
        record = _row_to_dict(header_values, values)
        _reservoir_add(reservoir, (row_index, record), seen, sample_size, rng)

    return [record for _row_number, record in sorted(reservoir, key=lambda item: item[0])]


def _row_to_dict(headers: list, values: list) -> dict:
    row = {}
    for index, header in enumerate(headers):
        header = frappe.as_unicode(header or "").strip()
        if not header:
            header = f"column_{index + 1}"
        row[header] = values[index] if index < len(values) else None
    return row


def _reservoir_add(reservoir: list, item, seen: int, sample_size: int, rng: random.Random) -> None:
    if sample_size <= 0:
        return
    if len(reservoir) < sample_size:
        reservoir.append(item)
        return
    replace_at = rng.randint(1, seen)
    if replace_at <= sample_size:
        reservoir[replace_at - 1] = item


def _has_values(values: list) -> bool:
    return any(value not in (None, "") for value in values)


def _use_random_sample(sample_strategy: str | None, max_sample_rows: int) -> bool:
    return (sample_strategy or "").lower() == "random" and max_sample_rows > 0


def _sheet_sample_seed(file_path: str, sheet_name: str, sample_seed: str | None) -> str:
    return f"{sample_seed or Path(file_path).name}::{sheet_name}"


def _to_jsonable(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if value is None:
        return None
    return value


def _looks_numeric(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except Exception:
        return False
