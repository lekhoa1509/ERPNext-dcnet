"""Workbook serialization and placeholder rendering for Excel print templates."""

from __future__ import annotations

import copy
import json
import re
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter, range_boundaries

from vn_accounting.excel_printing.constants import (
    MAX_COLUMNS,
    MAX_ROWS,
    MAX_SHEETS,
    WORKBOOK_SCHEMA_VERSION,
)

PLACEHOLDER_RE = re.compile(r"{{\s*([a-zA-Z_][\w.]*)\s*}}")
CELL_REF_RE = re.compile(r"\b([A-Z]{1,3}[1-9]\d*)\b")
SUM_RE = re.compile(r"SUM\(\s*([A-Z]{1,3}[1-9]\d*):([A-Z]{1,3}[1-9]\d*)\s*\)", re.I)


def workbook_to_state(workbook) -> dict[str, Any]:
    """Serialize an openpyxl workbook into the canonical builder state."""
    if len(workbook.worksheets) > MAX_SHEETS:
        raise ValueError(f"Workbook exceeds the {MAX_SHEETS}-sheet limit")

    sheets = []
    for worksheet in workbook.worksheets:
        max_row = min(max(worksheet.max_row, 1), MAX_ROWS)
        max_column = min(max(worksheet.max_column, 1), MAX_COLUMNS)
        cells = {}
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, max_col=max_column):
            for cell in row:
                if cell.value is None and not cell.has_style:
                    continue
                cells[cell.coordinate] = {
                    "value": _json_value(cell.value),
                    "data_type": cell.data_type,
                    "style": _serialize_style(cell),
                }

        sheets.append(
            {
                "title": worksheet.title,
                "max_row": max_row,
                "max_column": max_column,
                "cells": cells,
                "merged_cells": [str(item) for item in worksheet.merged_cells.ranges],
                "row_heights": {
                    str(index): dimension.height
                    for index, dimension in worksheet.row_dimensions.items()
                    if dimension.height is not None and index <= MAX_ROWS
                },
                "column_widths": {
                    key: dimension.width
                    for key, dimension in worksheet.column_dimensions.items()
                    if dimension.width is not None
                },
                "print_area": str(worksheet.print_area) if worksheet.print_area else "",
                "orientation": worksheet.page_setup.orientation or "portrait",
                "paper_size": worksheet.page_setup.paperSize or "9",
                "freeze_panes": str(worksheet.freeze_panes or ""),
                "show_grid_lines": bool(worksheet.sheet_view.showGridLines),
                "preview_bounds": _preview_bounds(worksheet),
            }
        )

    return {
        "version": WORKBOOK_SCHEMA_VERSION,
        "active_sheet": workbook.index(workbook.active),
        "sheets": sheets,
    }


def state_to_workbook(state: dict[str, Any]) -> Workbook:
    """Build an openpyxl workbook from canonical builder state."""
    state = validate_state(state)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_state in state["sheets"]:
        worksheet = workbook.create_sheet(sheet_state["title"][:31])
        worksheet.sheet_view.showGridLines = bool(sheet_state.get("show_grid_lines", False))
        worksheet.page_setup.orientation = sheet_state.get("orientation") or "portrait"
        worksheet.page_setup.paperSize = sheet_state.get("paper_size") or "9"
        if sheet_state.get("freeze_panes"):
            worksheet.freeze_panes = sheet_state["freeze_panes"]

        for column, width in sheet_state.get("column_widths", {}).items():
            worksheet.column_dimensions[column].width = float(width)
        for row, height in sheet_state.get("row_heights", {}).items():
            worksheet.row_dimensions[int(row)].height = float(height)

        for coordinate, cell_state in sheet_state.get("cells", {}).items():
            cell = worksheet[coordinate]
            cell.value = _restore_value(cell_state.get("value"), cell_state.get("data_type"))
            _apply_style(cell, cell_state.get("style") or {})

        for merged_range in sheet_state.get("merged_cells", []):
            if _valid_merge(merged_range, worksheet):
                worksheet.merge_cells(merged_range)
        if sheet_state.get("print_area"):
            worksheet.print_area = sheet_state["print_area"]

    workbook.active = min(int(state.get("active_sheet", 0)), len(workbook.worksheets) - 1)
    return workbook


def validate_state(state: dict[str, Any] | str) -> dict[str, Any]:
    """Validate and normalize workbook state received from the browser."""
    if isinstance(state, str):
        state = json.loads(state)
    if not isinstance(state, dict) or not isinstance(state.get("sheets"), list):
        raise ValueError("Invalid workbook state")
    if not state["sheets"] or len(state["sheets"]) > MAX_SHEETS:
        raise ValueError("Workbook must contain between 1 and 20 sheets")

    normalized = copy.deepcopy(state)
    normalized["version"] = WORKBOOK_SCHEMA_VERSION
    for sheet in normalized["sheets"]:
        if not isinstance(sheet, dict) or not sheet.get("title"):
            raise ValueError("Every sheet requires a title")
        sheet["title"] = str(sheet["title"])[:31]
        sheet["max_row"] = min(max(int(sheet.get("max_row") or 1), 1), MAX_ROWS)
        sheet["max_column"] = min(max(int(sheet.get("max_column") or 1), 1), MAX_COLUMNS)
        if not isinstance(sheet.get("cells", {}), dict):
            raise ValueError("Sheet cells must be an object")
        if len(sheet["cells"]) > MAX_ROWS * MAX_COLUMNS:
            raise ValueError("Sheet contains too many cells")
        for coordinate, cell in sheet["cells"].items():
            match = re.fullmatch(r"([A-Z]{1,3})([1-9]\d*)", coordinate)
            if not match:
                raise ValueError(f"Invalid cell coordinate: {coordinate}")
            column = _column_index(match.group(1))
            row = int(match.group(2))
            if column > MAX_COLUMNS or row > MAX_ROWS:
                raise ValueError(f"Cell coordinate exceeds builder limits: {coordinate}")
            if not isinstance(cell, dict):
                raise ValueError(f"Invalid cell payload: {coordinate}")
            sheet["max_row"] = max(sheet["max_row"], row)
            sheet["max_column"] = max(sheet["max_column"], column)
        merged_cells = sheet.setdefault("merged_cells", [])
        if not isinstance(merged_cells, list) or len(merged_cells) > MAX_ROWS:
            raise ValueError("Invalid merged cell ranges")
        for merged_range in merged_cells:
            try:
                min_column, min_row, max_column, max_row = range_boundaries(str(merged_range).replace("$", ""))
            except ValueError as error:
                raise ValueError(f"Invalid merged cell range: {merged_range}") from error
            if min_column < 1 or min_row < 1 or max_column > MAX_COLUMNS or max_row > MAX_ROWS:
                raise ValueError(f"Merged cell range exceeds builder limits: {merged_range}")
    return normalized


def render_state(state: dict[str, Any] | str, document: Any) -> dict[str, Any]:
    """Fill document placeholders and expand child-table template rows."""
    workbook = state_to_workbook(validate_state(state))
    context = _as_mapping(document)
    for worksheet in workbook.worksheets:
        _expand_child_rows(worksheet, context)
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = replace_placeholders(cell.value, context)
    return workbook_to_state(workbook)


def replace_placeholders(value: str, context: dict[str, Any], row_context: dict | None = None) -> Any:
    """Replace placeholders in a string while retaining native types for full-cell tokens."""
    matches = list(PLACEHOLDER_RE.finditer(value))
    if not matches:
        return value

    if len(matches) == 1 and matches[0].span() == (0, len(value)):
        resolved = _resolve_path(matches[0].group(1), context, row_context)
        return "" if resolved is None else resolved

    output = value
    for match in reversed(matches):
        resolved = _resolve_path(match.group(1), context, row_context)
        replacement = "" if resolved is None else str(resolved)
        output = output[: match.start()] + replacement + output[match.end() :]
    return output


def evaluate_preview_formulas(state: dict[str, Any]) -> dict[str, Any]:
    """Resolve basic arithmetic and SUM formulas for browser print preview only."""
    result = copy.deepcopy(state)
    for sheet in result.get("sheets", []):
        cells = sheet.get("cells", {})
        for _ in range(4):
            changed = False
            for coordinate, cell in cells.items():
                formula = cell.get("value")
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                evaluated = _evaluate_formula(formula[1:], cells)
                if evaluated is not None:
                    cell["display_value"] = evaluated
                    changed = True
            if not changed:
                break
    return result


def _expand_child_rows(worksheet, context: dict[str, Any]) -> None:
    template_rows = []
    for row_number in range(1, worksheet.max_row + 1):
        paths = []
        for cell in worksheet[row_number]:
            if isinstance(cell.value, str):
                paths.extend(match.group(1) for match in PLACEHOLDER_RE.finditer(cell.value))
        prefixes = {path.split(".", 1)[0] for path in paths if "." in path}
        child_prefix = next(
            (prefix for prefix in prefixes if isinstance(context.get(prefix), list)),
            None,
        )
        if child_prefix:
            template_rows.append((row_number, child_prefix))

    for row_number, child_prefix in reversed(template_rows):
        rows = context.get(child_prefix) or []
        if not rows:
            for cell in worksheet[row_number]:
                if isinstance(cell.value, str):
                    cell.value = replace_placeholders(cell.value, context, {})
            continue

        extra_count = len(rows) - 1
        original_merges = [str(item) for item in worksheet.merged_cells.ranges]
        row_height = worksheet.row_dimensions[row_number].height
        if extra_count:
            worksheet.insert_rows(row_number + 1, extra_count)
            _shift_merges_for_insert(worksheet, original_merges, row_number, extra_count)

        source_cells = [
            {
                "column": cell.column,
                "value": cell.value,
                "style": copy.copy(cell._style),
                "number_format": cell.number_format,
                "protection": copy.copy(cell.protection),
                "alignment": copy.copy(cell.alignment),
            }
            for cell in worksheet[row_number]
        ]
        for index, child_row in enumerate(rows):
            target_row = row_number + index
            if row_height is not None:
                worksheet.row_dimensions[target_row].height = row_height
            for source in source_cells:
                target = worksheet.cell(target_row, source["column"])
                if index:
                    target._style = copy.copy(source["style"])
                    target.number_format = source["number_format"]
                    target.protection = copy.copy(source["protection"])
                    target.alignment = copy.copy(source["alignment"])
                if isinstance(source["value"], str):
                    target.value = replace_placeholders(source["value"], context, _as_mapping(child_row))
                elif index:
                    target.value = source["value"]


def _shift_merges_for_insert(worksheet, ranges: list[str], row_number: int, count: int) -> None:
    for merged in list(worksheet.merged_cells.ranges):
        worksheet.unmerge_cells(str(merged))
    for merged in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(merged)
        if min_row > row_number:
            min_row += count
            max_row += count
        elif min_row == row_number and max_row == row_number:
            for offset in range(count + 1):
                worksheet.merge_cells(
                    start_row=row_number + offset,
                    start_column=min_col,
                    end_row=row_number + offset,
                    end_column=max_col,
                )
            continue
        worksheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )


def _resolve_path(path: str, context: dict[str, Any], row_context: dict | None) -> Any:
    parts = path.split(".")
    if parts[0] in {"doc", "document"}:
        parts = parts[1:]
        current: Any = context
    elif row_context is not None and (parts[0] in context and isinstance(context.get(parts[0]), list)):
        parts = parts[1:]
        current = row_context
    elif row_context is not None and parts[0] in {"item", "row"}:
        parts = parts[1:]
        current = row_context
    else:
        current = context

    for part in parts:
        if isinstance(current, dict):
            current = current.get(part)
        else:
            current = getattr(current, part, None)
        if current is None:
            break
    return current


def _as_mapping(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    if hasattr(value, "as_dict"):
        return value.as_dict()
    return dict(vars(value)) if hasattr(value, "__dict__") else {}


def _evaluate_formula(expression: str, cells: dict[str, dict]) -> float | int | None:
    def cell_value(coordinate: str) -> float:
        cell = cells.get(coordinate, {})
        value = cell.get("display_value", cell.get("value", 0))
        return float(value) if isinstance(value, (int, float, Decimal)) else 0.0

    def sum_value(match) -> str:
        start_col, start_row, end_col, end_row = range_boundaries(f"{match.group(1)}:{match.group(2)}")
        total = 0.0
        for row in range(start_row, end_row + 1):
            for column in range(start_col, end_col + 1):
                total += cell_value(f"{get_column_letter(column)}{row}")
        return str(total)

    expression = SUM_RE.sub(sum_value, expression.upper())
    expression = CELL_REF_RE.sub(lambda match: str(cell_value(match.group(1))), expression)
    if not re.fullmatch(r"[\d\s.+\-*/()]+", expression):
        return None
    try:
        result = eval(expression, {"__builtins__": {}}, {})  # noqa: S307 - grammar restricted above
    except (ArithmeticError, SyntaxError, TypeError, ValueError):
        return None
    return result if isinstance(result, (int, float)) else None


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return {"__type": value.__class__.__name__, "value": value.isoformat()}
    if isinstance(value, Decimal):
        return float(value)
    return value


def _restore_value(value: Any, data_type: str | None) -> Any:
    if isinstance(value, dict) and value.get("__type"):
        kind = value["__type"]
        parser = {"datetime": datetime.fromisoformat, "date": date.fromisoformat, "time": time.fromisoformat}.get(kind)
        return parser(value["value"]) if parser else value["value"]
    return value


def _serialize_style(cell) -> dict[str, Any]:
    return {
        "font": {
            "name": cell.font.name,
            "size": cell.font.sz,
            "bold": bool(cell.font.b),
            "italic": bool(cell.font.i),
            "underline": cell.font.u,
            "strike": bool(cell.font.strike),
            "color": _serialize_color(cell.font.color),
        },
        "fill": {
            "fill_type": cell.fill.fill_type,
            "fg_color": _serialize_color(cell.fill.fgColor),
            "bg_color": _serialize_color(cell.fill.bgColor),
        },
        "border": {
            side: _serialize_side(getattr(cell.border, side))
            for side in ("left", "right", "top", "bottom")
        },
        "alignment": {
            "horizontal": cell.alignment.horizontal,
            "vertical": cell.alignment.vertical,
            "wrap_text": bool(cell.alignment.wrap_text),
            "shrink_to_fit": bool(cell.alignment.shrink_to_fit),
            "text_rotation": cell.alignment.text_rotation,
            "indent": cell.alignment.indent,
        },
        "number_format": cell.number_format,
        "protection": {"locked": cell.protection.locked, "hidden": cell.protection.hidden},
    }


def _apply_style(cell, style: dict[str, Any]) -> None:
    font = style.get("font", {})
    cell.font = Font(
        name=font.get("name"), size=font.get("size"), bold=font.get("bold", False),
        italic=font.get("italic", False), underline=font.get("underline"),
        strike=font.get("strike", False), color=_restore_color(font.get("color")),
    )
    fill = style.get("fill", {})
    cell.fill = PatternFill(
        fill_type=fill.get("fill_type"),
        fgColor=_restore_color(fill.get("fg_color")) or Color(rgb="00000000"),
        bgColor=_restore_color(fill.get("bg_color")) or Color(rgb="00000000"),
    )
    border = style.get("border", {})
    cell.border = Border(**{name: _restore_side(border.get(name)) for name in ("left", "right", "top", "bottom")})
    alignment = style.get("alignment", {})
    cell.alignment = Alignment(
        horizontal=alignment.get("horizontal"), vertical=alignment.get("vertical"),
        wrap_text=alignment.get("wrap_text", False),
        shrink_to_fit=alignment.get("shrink_to_fit", False),
        text_rotation=alignment.get("text_rotation", 0), indent=alignment.get("indent", 0),
    )
    cell.number_format = style.get("number_format") or "General"
    protection = style.get("protection", {})
    cell.protection = Protection(locked=protection.get("locked", True), hidden=protection.get("hidden", False))


def _serialize_color(color) -> dict[str, Any] | None:
    if color is None:
        return None
    return {
        "type": color.type,
        "rgb": color.rgb if color.type == "rgb" else None,
        "indexed": color.indexed if color.type == "indexed" else None,
        "theme": color.theme if color.type == "theme" else None,
        "tint": color.tint,
    }


def _restore_color(value: dict[str, Any] | None) -> Color | None:
    if not value:
        return None
    kwargs = {"tint": value.get("tint", 0)}
    if value.get("type") == "rgb" and value.get("rgb"):
        kwargs["rgb"] = value["rgb"]
    elif value.get("type") == "indexed" and value.get("indexed") is not None:
        kwargs["indexed"] = value["indexed"]
    elif value.get("type") == "theme" and value.get("theme") is not None:
        kwargs["theme"] = value["theme"]
    return Color(**kwargs)


def _serialize_side(side) -> dict[str, Any]:
    if side is None:
        return {"style": None, "color": None}
    return {"style": side.style, "color": _serialize_color(side.color)}


def _restore_side(value: dict[str, Any] | None) -> Side:
    value = value or {}
    return Side(style=value.get("style"), color=_restore_color(value.get("color")))


def _column_index(name: str) -> int:
    result = 0
    for character in name:
        result = result * 26 + ord(character) - 64
    return result


def _valid_merge(merged_range: str, worksheet) -> bool:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(merged_range)
    except ValueError:
        return False
    non_empty = 0
    for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        non_empty += sum(cell.value not in (None, "") for cell in row)
    return non_empty <= 1


def _preview_bounds(worksheet) -> dict[str, int]:
    """Return a compact print-preview area, preferring an explicit Excel print area."""
    print_ranges = re.findall(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)", str(worksheet.print_area or ""))
    if print_ranges:
        ranges = [
            range_boundaries(f"{start_column}{start_row}:{end_column}{end_row}")
            for start_column, start_row, end_column, end_row in print_ranges
        ]
        return {
            "min_column": min(item[0] for item in ranges),
            "min_row": min(item[1] for item in ranges),
            "max_column": min(max(item[2] for item in ranges), MAX_COLUMNS),
            "max_row": min(max(item[3] for item in ranges), MAX_ROWS),
        }

    coordinates = [cell.coordinate for row in worksheet.iter_rows() for cell in row if cell.value not in (None, "")]
    for merged in worksheet.merged_cells.ranges:
        min_column, min_row, max_column, max_row = range_boundaries(str(merged))
        coordinates.extend((f"{get_column_letter(min_column)}{min_row}", f"{get_column_letter(max_column)}{max_row}"))
    if not coordinates:
        return {"min_column": 1, "min_row": 1, "max_column": 1, "max_row": 1}

    boundaries = [range_boundaries(f"{coordinate}:{coordinate}") for coordinate in coordinates]
    return {
        "min_column": min(item[0] for item in boundaries),
        "min_row": min(item[1] for item in boundaries),
        "max_column": min(max(item[2] for item in boundaries), MAX_COLUMNS),
        "max_row": min(max(item[3] for item in boundaries), MAX_ROWS),
    }


def load_xlsx(path: str):
    """Load an XLSX/XLSM workbook with formulas and formatting preserved."""
    return load_workbook(path, data_only=False, keep_vba=path.lower().endswith(".xlsm"))
