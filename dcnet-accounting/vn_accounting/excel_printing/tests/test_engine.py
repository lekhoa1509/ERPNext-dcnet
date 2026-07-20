"""Deterministic tests for workbook state and placeholder rendering."""

from unittest import TestCase

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from vn_accounting.excel_printing.engine import (
    evaluate_preview_formulas,
    render_state,
    state_to_workbook,
    validate_state,
    workbook_to_state,
)


class TestExcelPrintEngine(TestCase):
    def test_workbook_style_roundtrip(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Phiếu thu"
        worksheet["A1"] = "CÔNG TY DCNET"
        worksheet["A1"].font = Font(name="Times New Roman", size=14, bold=True, color="FF123456")
        worksheet["A1"].fill = PatternFill("solid", fgColor="FFE0F2FE")
        worksheet["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
        worksheet["A1"].border = Border(bottom=Side(style="thin", color="FF000000"))
        worksheet.merge_cells("A1:C1")
        worksheet.column_dimensions["A"].width = 24
        worksheet.row_dimensions[1].height = 28

        restored = state_to_workbook(workbook_to_state(workbook))
        cell = restored["Phiếu thu"]["A1"]
        self.assertEqual(cell.value, "CÔNG TY DCNET")
        self.assertEqual(cell.font.name, "Times New Roman")
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.alignment.horizontal, "center")
        self.assertIn("A1:C1", [str(item) for item in restored["Phiếu thu"].merged_cells.ranges])
        self.assertEqual(restored["Phiếu thu"].column_dimensions["A"].width, 24)

    def test_document_and_child_placeholders_expand(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "Khách hàng: {{customer_name}}"
        worksheet["A3"] = "{{items.item_code}}"
        worksheet["B3"] = "{{items.item_name}}"
        worksheet["C3"] = "{{items.qty}}"
        worksheet["A4"] = "Tổng cộng"

        rendered = render_state(
            workbook_to_state(workbook),
            {
                "customer_name": "Công ty Thăng Long",
                "items": [
                    {"item_code": "SP-001", "item_name": "Gậy golf", "qty": 1},
                    {"item_code": "SP-002", "item_name": "Bóng golf", "qty": 2},
                ],
            },
        )
        restored = state_to_workbook(rendered).active
        self.assertEqual(restored["A1"].value, "Khách hàng: Công ty Thăng Long")
        self.assertEqual(restored["A3"].value, "SP-001")
        self.assertEqual(restored["A4"].value, "SP-002")
        self.assertEqual(restored["A5"].value, "Tổng cộng")

    def test_formula_preview_supports_arithmetic_and_sum(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = 10
        worksheet["A2"] = 20
        worksheet["A3"] = "=SUM(A1:A2)"
        worksheet["A4"] = "=A3*2"
        state = evaluate_preview_formulas(workbook_to_state(workbook))
        cells = state["sheets"][0]["cells"]
        self.assertEqual(cells["A3"]["display_value"], 30)
        self.assertEqual(cells["A4"]["display_value"], 60)

    def test_preview_bounds_use_print_area(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["A1"] = "Ngoài vùng in"
        worksheet["C3"] = "Bắt đầu"
        worksheet["D4"] = "Kết thúc"
        worksheet.print_area = "C3:D4"

        bounds = workbook_to_state(workbook)["sheets"][0]["preview_bounds"]

        self.assertEqual(bounds, {"min_column": 3, "min_row": 3, "max_column": 4, "max_row": 4})

    def test_state_rejects_cells_outside_builder_limits(self):
        state = {
            "sheets": [{
                "title": "Sheet1",
                "cells": {"XFD1048576": {"value": "too large"}},
            }],
        }

        with self.assertRaises(ValueError):
            validate_state(state)
