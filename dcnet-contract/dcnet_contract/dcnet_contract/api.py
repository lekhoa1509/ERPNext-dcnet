"""Public API for DCNet Contract Template Builder."""

import os
import re

import frappe
from frappe import _
from frappe.utils import getdate, today

# ---------------------------------------------------------------------------
# Field groups for the template builder sidebar
# ---------------------------------------------------------------------------
_FIELD_GROUPS = [
	("Bên A (Khách hàng)", [
		"customer_name", "customer_address", "customer_phone", "customer_fax",
		"customer_email", "customer_tax_id", "customer_representative",
		"customer_representative_title", "customer_bank_account",
		"customer_id_number", "customer_id_date", "customer_id_place", "customer_dob",
	]),
	("Bên B (Công ty)", [
		"company_name_b", "company_address_b", "company_phone_b", "company_fax_b",
		"company_tax_id_b", "company_representative_b", "company_representative_title_b",
		"company_bank_account_b", "company_bank_b",
	]),
	("Hợp đồng", [
		"contract_number", "contract_date", "acceptance_date", "end_date",
		"package_term", "service_type", "package_name",
	]),
	("Địa điểm & Kỹ thuật", [
		"installation_address", "bandwidth", "point_a", "point_b", "sla_restore_hours",
	]),
	("Phí & Thanh toán", [
		"monthly_fee", "monthly_fee_vat", "monthly_fee_total", "monthly_fee_words",
		"setup_fee", "setup_fee_vat", "setup_fee_total", "setup_fee_words",
	]),
	("Hạng mục (vòng lặp trong bảng)", [
		"item_stt", "item_label", "item_qty", "item_uom", "item_price", "item_amount",
	]),
]


@frappe.whitelist()
def get_contract_commission_states(contract: str) -> dict:
	"""Return the highest-priority commission state for each billing period."""
	if not contract:
		return {}
	if not frappe.db.exists("DCNet Contract", contract):
		frappe.throw(_("Contract not found: {0}").format(contract), frappe.DoesNotExistError)
	if not frappe.has_permission("DCNet Contract", "read", contract):
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	if not frappe.db.exists("DocType", "Phuong An Kinh Doanh"):
		return {}

	pakds = frappe.get_all(
		"Phuong An Kinh Doanh",
		filters={"contract_ref": contract},
		pluck="name",
	)
	if not pakds:
		return {}

	rows = frappe.db.sql(
		"""
		SELECT billing_schedule_idx, state
		FROM `tabPAKD Commission Line`
		WHERE parent IN %(pakds)s
		""",
		{"pakds": pakds},
		as_dict=True,
	)
	priority = {"Posted": 3, "Pending": 2, "Skipped": 1, "Cancelled": 0}
	by_period: dict[int, str] = {}
	for row in rows:
		idx = row.billing_schedule_idx
		current = by_period.get(idx)
		if current is None or priority.get(row.state, -1) > priority.get(current, -1):
			by_period[idx] = row.state
	return by_period


@frappe.whitelist()
def get_contract_fields():
	"""Return grouped placeholder fields for the template builder sidebar."""
	from dcnet_contract.dcnet_contract.utils.docx_generator import PLACEHOLDER_MAP

	result = []
	for group_name, keys in _FIELD_GROUPS:
		fields = []
		for key in keys:
			_mapped, label = PLACEHOLDER_MAP.get(key, ("", key))
			fields.append({"key": key, "label": label or key})
		result.append({"group": group_name, "fields": fields})
	return result


@frappe.whitelist()
def get_print_templates():
	"""Return all contract templates for the print dialog."""
	return frappe.get_all(
		"DCNet Contract Template",
		fields=["name", "template_name", "template_category", "service_type", "status", "template_file"],
		order_by="template_name asc",
	)


@frappe.whitelist()
def docx_to_html(file_url):
	"""Convert an uploaded .docx file to styled HTML using python-docx.

	Preserves paragraph alignment, bold/italic/underline, font name & size,
	text color, and table structure as inline styles.
	"""
	try:
		file_doc = frappe.get_doc("File", {"file_url": file_url})
		file_path = file_doc.get_full_path()

		if not file_path or not os.path.exists(file_path):
			frappe.throw(f"Không tìm thấy file: {file_url}")

		html = _docx_to_styled_html(file_path)
		return {"html": html}

	except Exception as e:
		frappe.log_error(frappe.get_traceback(), "dcnet_contract docx_to_html error")
		frappe.throw(f"Lỗi chuyển đổi file Word: {str(e)}")


def _docx_to_styled_html(file_path: str) -> str:
	"""Convert .docx to HTML with full inline style preservation via python-docx."""
	from docx import Document as DocxDocument
	from docx.enum.text import WD_ALIGN_PARAGRAPH
	from docx.oxml.ns import qn
	from docx.text.paragraph import Paragraph
	from docx.table import Table

	doc = DocxDocument(file_path)

	_ALIGN_CSS = {
		WD_ALIGN_PARAGRAPH.CENTER: "center",
		WD_ALIGN_PARAGRAPH.RIGHT: "right",
		WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
		WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
		WD_ALIGN_PARAGRAPH.LEFT: "left",
	}

	def _esc(s):
		return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

	def _para_align(para):
		align = para.alignment
		if align is None:
			align = para.paragraph_format.alignment
		if align is None and para.style and para.style.paragraph_format:
			align = para.style.paragraph_format.alignment
		return _ALIGN_CSS.get(align, "")

	def _run_to_html(run):
		text = run.text
		if not text:
			return ""
		escaped = _esc(text)
		styles = []
		if run.bold:
			styles.append("font-weight:bold")
		if run.italic:
			styles.append("font-style:italic")
		if run.underline:
			styles.append("text-decoration:underline")
		elif run.font.strike:
			styles.append("text-decoration:line-through")
		if run.font.name:
			styles.append(f"font-family:'{run.font.name}'")
		if run.font.size:
			styles.append(f"font-size:{round(run.font.size.pt, 1)}pt")
		try:
			if run.font.color and run.font.color.type is not None:
				rgb = str(run.font.color.rgb)
				if rgb not in ("000000", "auto"):
					styles.append(f"color:#{rgb}")
		except Exception:
			pass
		if styles:
			return f'<span style="{";".join(styles)}">{escaped}</span>'
		return escaped

	def _para_to_html(para):
		style_name = (para.style.name or "") if para.style else ""

		tag = "p"
		if any(n in style_name for n in ("Heading 1", "Title")):
			tag = "h1"
		elif any(n in style_name for n in ("Heading 2", "Subtitle")):
			tag = "h2"
		elif "Heading 3" in style_name:
			tag = "h3"

		align = _para_align(para)

		inner_parts = []
		for elem in para._element:
			local = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
			if local == "r":
				from docx.text.run import Run
				inner_parts.append(_run_to_html(Run(elem, para)))
			elif local == "hyperlink":
				for r_elem in elem.findall(qn("w:r")):
					from docx.text.run import Run
					inner_parts.append(_run_to_html(Run(r_elem, para)))

		inner = "".join(inner_parts)
		if not inner.strip():
			return "<p><br></p>"

		style_attr = f' style="text-align:{align}"' if align and align != "left" else ""
		return f"<{tag}{style_attr}>{inner}</{tag}>"

	def _cell_to_html(cell):
		return "".join(_para_to_html(p) for p in cell.paragraphs)

	def _table_to_html(table):
		rows = []
		for row in table.rows:
			cells_html = ""
			all_cells = row.cells
			i = 0
			while i < len(all_cells):
				cell = all_cells[i]
				# Count how many consecutive cells share the same XML element (merged)
				colspan = 1
				while i + colspan < len(all_cells) and all_cells[i + colspan]._tc is cell._tc:
					colspan += 1
				cs_attr = f' colspan="{colspan}"' if colspan > 1 else ""
				cells_html += (
					f'<td{cs_attr} style="padding:4px 8px;border:1px solid #999;vertical-align:top">'
					f'{_cell_to_html(cell)}</td>'
				)
				i += colspan
			rows.append(f"<tr>{cells_html}</tr>")
		return (
			'<table style="width:100%;border-collapse:collapse;margin:8px 0">'
			+ "".join(rows)
			+ "</table>"
		)

	parts = []
	for child in doc.element.body:
		local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
		if local == "p":
			parts.append(_para_to_html(Paragraph(child, doc)))
		elif local == "tbl":
			parts.append(_table_to_html(Table(child, doc)))

	return "".join(parts)


@frappe.whitelist()
def save_template_from_builder(template_name_key, display_name, html_content,
								template_category="", service_type="", docx_file_url=""):
	"""Create or update a DCNet Contract Template from the builder page."""
	from dcnet_contract.dcnet_contract.utils.docx_generator import PLACEHOLDER_MAP

	if template_name_key and frappe.db.exists("DCNet Contract Template", template_name_key):
		tpl = frappe.get_doc("DCNet Contract Template", template_name_key)
	else:
		tpl = frappe.new_doc("DCNet Contract Template")
		tpl.template_name = display_name

	tpl.template_html = html_content
	tpl.template_category = template_category or ""
	tpl.service_type = service_type or ""
	if docx_file_url:
		tpl.template_file = docx_file_url

	# Auto-discover placeholders from HTML
	found_keys = set(re.findall(r"\{\{(\w+)\}\}", html_content))
	tpl.set("placeholders", [])
	for key in sorted(found_keys):
		_mapped, desc = PLACEHOLDER_MAP.get(key, ("", ""))
		tpl.append("placeholders", {
			"placeholder_key": key,
			"mapped_field": _mapped,
			"description": desc,
		})

	tpl.save(ignore_permissions=True)
	frappe.db.commit()
	return {"name": tpl.name, "template_name": tpl.template_name}


@frappe.whitelist()
def delete_print_template(template_name):
	"""Delete a contract template."""
	frappe.delete_doc("DCNet Contract Template", template_name, ignore_permissions=True)
	frappe.db.commit()
	return {"status": "ok"}


@frappe.whitelist()
def get_template_html(template_name):
	"""Return the template_html and metadata for the builder."""
	tpl = frappe.get_doc("DCNet Contract Template", template_name)
	return {
		"name": tpl.name,
		"template_name": tpl.template_name,
		"template_category": tpl.template_category or "",
		"service_type": tpl.service_type or "",
		"template_html": tpl.template_html or tpl.boilerplate_text or "",
		"template_file": tpl.template_file or "",
	}


@frappe.whitelist()
def debug_templates():
	"""Debug helper - list all templates and their content."""
	tpls = frappe.get_all("DCNet Contract Template", fields=["name", "template_name", "template_html", "template_file", "boilerplate_text", "service_type"])
	result = []
	for t in tpls:
		result.append({
			"name": t.name,
			"template_name": t.template_name,
			"service_type": t.service_type,
			"has_template_html": bool(t.template_html),
			"template_html_len": len(t.template_html) if t.template_html else 0,
			"has_template_file": bool(t.template_file),
			"has_boilerplate": bool(t.boilerplate_text),
			"template_html_preview": (t.template_html or "")[:100] if t.template_html else "",
		})
	return result


@frappe.whitelist()
def preview_template_with_contract(contract_name, template_name):
	"""Render template HTML with real data from a DCNet Contract for preview."""
	import mammoth
	from dcnet_contract.dcnet_contract.utils.docx_generator import (
		_build_placeholder_values,
		_build_items_list,
		_contains_item_placeholders,
		_fill_html_template,
	)

	# Get template HTML
	tpl = frappe.get_doc("DCNet Contract Template", template_name)
	html = tpl.template_html or tpl.boilerplate_text or ""

	has_no_placeholders = not re.search(r"\{\{\w+\}\}", html)

	# If no HTML but has DOCX file, convert it
	if not html and tpl.template_file:
		file_doc = frappe.get_doc("File", {"file_url": tpl.template_file})
		file_path = file_doc.get_full_path()
		with open(file_path, "rb") as f:
			result = mammoth.convert_to_html(
				f,
				style_map="""
					r[style-name='Heading 1'] => h1
					r[style-name='Heading 2'] => h2
					r[style-name='Heading 3'] => h3
					p[style-name='Normal'] => p
				""",
			)
		html = _preserve_word_styles(result.value)

	if not html:
		return {"html": "<p>Mẫu này chưa có nội dung.</p>", "template_name": tpl.template_name}

	# Build placeholder values from real contract data
	values = _build_placeholder_values(contract_name)

	# Build items list for table cloning
	items = _build_items_list(contract_name)
	html = _fill_html_template(html, values, items)

	# Replace remaining placeholders with contract values
	for key, val in values.items():
		html = html.replace("{{" + key + "}}", str(val))

	# Highlight any remaining unknown placeholders
	remaining = re.findall(r'\{\{(\w+)\}\}', html)
	if remaining:
		for ph in set(remaining):
			html = html.replace("{{" + ph + "}}", f'<span style="background:#fef08a;color:#92400e;padding:0 4px;border-radius:2px;font-size:11px;font-family:monospace">{{{{{ph}}}}}</span>')

	return {
		"html": html,
		"template_name": tpl.template_name,
		"contract_name": contract_name,
		"items_count": len(items),
		"no_placeholders": has_no_placeholders and not _contains_item_placeholders(html),
	}


@frappe.whitelist()
def preview_template_html(contract_name, template_html):
	"""Replace placeholders in raw template HTML with real contract data.

	Used by template builder preview when a contract context is available.
	"""
	from dcnet_contract.dcnet_contract.utils.docx_generator import (
		_build_placeholder_values,
		_build_items_list,
		_fill_html_template,
	)

	# Build real data from contract
	values = _build_placeholder_values(contract_name)
	items = _build_items_list(contract_name)

	html = _fill_html_template(template_html, values, items)

	# Replace all placeholders with real values
	for key, val in values.items():
		html = html.replace("{{" + key + "}}", str(val))

	# Highlight remaining unknown placeholders
	remaining = re.findall(r'\{\{(\w+)\}\}', html)
	if remaining:
		for ph in set(remaining):
			html = html.replace("{{" + ph + "}}", f'<span style="background:#fef08a;color:#92400e;padding:0 4px;border-radius:2px;font-size:11px;font-family:monospace">{{{{{ph}}}}}</span>')

	return {"html": html, "contract_name": contract_name}


@frappe.whitelist()
def get_summary_kpis(contract_name: str) -> dict:
	"""Return the KPI values and next action for a contract summary card."""
	if not contract_name or not frappe.db.exists("DCNet Contract", contract_name):
		frappe.throw(_("Contract not found: {0}").format(contract_name), frappe.DoesNotExistError)
	if not frappe.has_permission("DCNet Contract", "read", contract_name):
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	contract = frappe.get_cached_doc("DCNet Contract", contract_name)
	si_names = frappe.db.sql_list(
		"""SELECT DISTINCT sales_invoice
		   FROM `tabDCNet Contract Billing Schedule`
		   WHERE parent=%s AND sales_invoice IS NOT NULL AND sales_invoice != ''""",
		contract_name,
	)

	total_billed = 0.0
	total_collected = 0.0
	outstanding = 0.0
	if si_names:
		rows = frappe.db.sql(
			"""SELECT
				IFNULL(SUM(grand_total), 0) AS billed,
				IFNULL(SUM(grand_total - outstanding_amount), 0) AS collected,
				IFNULL(SUM(outstanding_amount), 0) AS outstanding
			   FROM `tabSales Invoice`
			   WHERE name IN %(names)s AND docstatus=1""",
			{"names": tuple(si_names)},
			as_dict=True,
		)
		if rows:
			total_billed = float(rows[0].billed or 0)
			total_collected = float(rows[0].collected or 0)
			outstanding = float(rows[0].outstanding or 0)

	overdue_count = frappe.db.count(
		"DCNet Contract Billing Schedule",
		{"parent": contract_name, "state": "Overdue"},
	)
	end_date = contract.end_date
	days_remaining = None
	if end_date:
		days_remaining = (getdate(end_date) - getdate(today())).days

	return {
		"status": contract.status,
		"contract_type": contract.contract_type,
		"service_type": contract.service_type,
		"customer": contract.customer,
		"customer_name": contract.customer_name,
		"sales_person_name": contract.sales_person_name,
		"branch": contract.branch,
		"grand_total": float(contract.grand_total or 0),
		"total_billed": total_billed,
		"total_collected": total_collected,
		"outstanding": outstanding,
		"overdue_count": overdue_count,
		"contract_date": contract.contract_date,
		"acceptance_date": contract.acceptance_date,
		"end_date": end_date,
		"days_remaining": days_remaining,
		"snapshot_date": contract.contract_date,
		"is_empty": not contract.items or (contract.grand_total or 0) == 0,
		"next_action": _contract_next_action(contract, overdue_count, days_remaining),
	}


def _contract_next_action(contract, overdue_count: int, days_remaining):
	"""Determine the highest-priority action for the contract summary card."""
	if contract.status == "Draft":
		return {
			"text": _("Hoàn thành thông tin → gửi duyệt"),
			"button_label": None,
			"button_action": None,
		}

	if contract.status == "Active":
		overdue_invoice = _find_overdue_invoice(contract.name)
		if overdue_invoice:
			return {
				"text": _("Theo dõi công nợ HĐ {0}: còn nợ {1}").format(
					overdue_invoice["name"],
					frappe.format_value(
						overdue_invoice["outstanding_amount"],
						{"fieldtype": "Currency"},
					),
				),
				"button_label": _("Mở phiếu"),
				"button_action": "navigate",
				"button_args": {"doctype": "Sales Invoice", "name": overdue_invoice["name"]},
			}

		if days_remaining is not None and days_remaining < 60:
			return {
				"text": _("Hợp đồng sắp hết hạn ({0} ngày) — chuẩn bị gia hạn").format(days_remaining),
				"button_label": _("Tạo HĐ gia hạn"),
				"button_action": "amend_contract",
				"button_args": {"contract": contract.name},
			}

		due_row = _find_due_billing_row(contract.name)
		if due_row:
			return {
				"text": _("Xuất hoá đơn kỳ {0} (đến hạn {1})").format(
					f"T{due_row['month_index']}",
					frappe.format_value(due_row["due_date"], {"fieldtype": "Date"}),
				),
				"button_label": _("Tạo hoá đơn"),
				"button_action": "create_invoice",
				"button_args": {"contract": contract.name, "billing_row": due_row["name"]},
			}

		next_row = _find_next_projected_billing_row(contract.name)
		if next_row:
			return {
				"text": _("Tự động chu kỳ tiếp theo: {0}").format(
					frappe.format_value(next_row["due_date"], {"fieldtype": "Date"})
				),
				"button_label": None,
				"button_action": None,
			}

		return {
			"text": _("Hợp đồng đang hoạt động — không có việc cần làm"),
			"button_label": None,
			"button_action": None,
		}

	if contract.status == "Suspended":
		return {
			"text": _("Hợp đồng đang tạm ngưng"),
			"button_label": _("Kích hoạt lại"),
			"button_action": "resume_contract",
			"button_args": {"name": contract.name},
		}
	if contract.status == "Expired":
		return {
			"text": _("Hợp đồng đã hết hạn — đóng hoặc gia hạn"),
			"button_label": _("Gia hạn"),
			"button_action": "amend_contract",
			"button_args": {"contract": contract.name},
		}
	if contract.status == "Cancelled":
		return {"text": _("Đã huỷ"), "button_label": None, "button_action": None}
	if contract.status == "Revised":
		return {
			"text": _("Đã sửa đổi sang phiên bản mới"),
			"button_label": None,
			"button_action": None,
		}
	return {"text": "", "button_label": None, "button_action": None}


def _find_overdue_invoice(contract_name: str):
	"""Return the oldest overdue submitted invoice for a contract."""
	rows = frappe.db.sql(
		"""SELECT si.name, si.outstanding_amount, si.due_date
		   FROM `tabSales Invoice` si
		   INNER JOIN `tabDCNet Contract Billing Schedule` bs
		     ON bs.sales_invoice = si.name
		   WHERE bs.parent=%s
		     AND si.docstatus=1
		     AND si.outstanding_amount > 0
		     AND si.due_date < %s
		   ORDER BY si.due_date ASC LIMIT 1""",
		(contract_name, today()),
		as_dict=True,
	)
	return rows[0] if rows else None


def _find_due_billing_row(contract_name: str):
	"""Return the oldest projected billing row that is due."""
	rows = frappe.db.sql(
		"""SELECT name, month_index, due_date
		   FROM `tabDCNet Contract Billing Schedule`
		   WHERE parent=%s AND state='Projected' AND due_date <= %s
		   ORDER BY due_date ASC LIMIT 1""",
		(contract_name, today()),
		as_dict=True,
	)
	return rows[0] if rows else None


def _find_next_projected_billing_row(contract_name: str):
	"""Return the next future projected billing row."""
	rows = frappe.db.sql(
		"""SELECT name, month_index, due_date
		   FROM `tabDCNet Contract Billing Schedule`
		   WHERE parent=%s AND state='Projected' AND due_date > %s
		   ORDER BY due_date ASC LIMIT 1""",
		(contract_name, today()),
		as_dict=True,
	)
	return rows[0] if rows else None
